import unittest
from pathlib import Path
from openpyxl import load_workbook
from relatorio_gerencial import build_report
from relatorio_modelo import MODEL_SHEETS, PRIMARY_SHEETS, TECHNICAL_SHEETS
from test_gerencial import fixture


class ModeloTest(unittest.TestCase):
    def setUp(self):
        self.wb=fixture()
        build_report(self.wb)

    def test_visible_tabs_order_and_original_data(self):
        self.assertEqual([s.title for s in self.wb if s.sheet_state=='visible'], [n for n in PRIMARY_SHEETS + TECHNICAL_SHEETS if n in self.wb.sheetnames])
        self.assertEqual(self.wb.active.title,'Dashboard')
        self.assertEqual(self.wb['Chat_ao_vivo'].max_row,9)
        self.assertEqual(self.wb['Transcrição dos Chats'].max_row,7)

    def test_headers_match_template(self):
        template=load_workbook(Path(__file__).with_name('modelo_relatorio.xlsx'))
        for name in MODEL_SHEETS:
            if name in {'Dashboard', 'Resumo da transmissão', 'Picos de 5 minutos'}:
                continue  # essas abas recebem rótulos mais claros na camada final
            row=2 if name in {'Participação','Metodologia e lacunas'} else 1
            source = 'Resumo por vídeo' if name == 'Resumo da transmissão' else name
            expected=[c.value for c in template[source][row]]
            actual=[self.wb[name].cell(row,i+1).value for i in range(len(expected))]
            self.assertEqual(expected,actual,name)
        self.assertEqual(self.wb['Resumo da transmissão']['A1'].value, 'Atividade Youtube')
        self.assertEqual(self.wb['Resumo da transmissão'].cell(1, self.wb['Resumo da transmissão'].max_column).value, 'Duração (h:mm:ss)')
        template.close()

    def test_reconciled_views_and_no_reference_data(self):
        self.assertEqual(self.wb['Dashboard']['A2'].value,'Relatório da transmissão')
        self.assertEqual(self.wb['Resumo da transmissão']['F2'].value,4)
        self.assertEqual(sum(self.wb['Resumo_Categorias'].cell(i,3).value for i in range(2,12)),4)
        self.assertEqual(self.wb['Participação']['B3'].value,2)
        self.assertEqual(self.wb['Resumo da transmissão']['N2'].value,0)
        self.assertFalse(any(c.data_type=='f' for name in MODEL_SHEETS for row in self.wb[name] for c in row))
        self.assertFalse(any('Painel 6' in str(c.value) for name in MODEL_SHEETS for row in self.wb[name] for c in row))


if __name__=='__main__':unittest.main(verbosity=2)
