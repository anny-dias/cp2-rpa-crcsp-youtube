#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations
import argparse, hashlib, json, logging, os, platform, re, shutil, subprocess, sys, time
from datetime import datetime, timezone
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from etl_core import *
from entrada_youtube import normalizar_link, solicitar_link
from integracao_web import validar_video_youtube
from relatorio_gerencial import adicionar_relatorio
from saida_segura import criar_temporario, publicar_excel
LOG=logging.getLogger("youtube_etl")

def setup(verbose=False): logging.basicConfig(level=logging.DEBUG if verbose else logging.INFO,format="%(asctime)s | %(levelname)s | %(message)s",datefmt="%H:%M:%S")
def run(cmd,cwd=None): return subprocess.run(cmd,cwd=str(cwd) if cwd else None,text=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,encoding="utf-8",errors="replace",check=False)
def retryable(t):
    t=(t or "").lower(); return any(x in t for x in ["http error 429","too many requests","http error 5","timed out","timeout","temporarily unavailable","connection reset","remote end closed"])
def run_retry(cmd,cwd,label,attempts,backoff):
    last=None
    for i in range(1,attempts+1):
        p=run(cmd,cwd); last=p
        if p.returncode==0:return p
        if i>=attempts or not retryable(p.stdout):return p
        delay=min(backoff*(2**(i-1)),120); LOG.warning("%s falhou %d/%d; nova tentativa em %ds",label,i,attempts,delay); time.sleep(delay)
    return last
def ya(browser=None):
    a=[sys.executable,"-m","yt_dlp","--retries","5","--fragment-retries","5","--retry-sleep","http:exp=5:60","--sleep-requests","1"]
    if browser:a += ["--cookies-from-browser",browser]
    return a
def video_id(url):
    return normalizar_link(url)[1]
def metadata(url,work,browser,attempts,backoff):
    p=run_retry(ya(browser)+["--skip-download","--no-playlist","--dump-single-json",url],work,"metadados",attempts,backoff)
    if p.returncode:raise RuntimeError(sanitize_log_detail(p.stdout))
    for line in reversed(p.stdout.splitlines()):
        line=line.strip()
        if line.startswith("{") and line.endswith("}"):return json.loads(line)
    raise RuntimeError("JSON de metadados não localizado")
def parse_json3(path):
    d=json.loads(path.read_text(encoding="utf-8",errors="replace"));rows=[]
    for e in d.get("events",[]):
        txt="".join(str(x.get("utf8","")) for x in (e.get("segs") or [])).replace("\n"," ").strip()
        if not txt:continue
        st=float(e.get("tStartMs",0))/1000;en=st+float(e.get("dDurationMs",0))/1000
        rows.append([st,en,seconds_to_hms(st),seconds_to_hms(en),txt,"legenda_youtube"])
    return pd.DataFrame(rows,columns=RAW_TRANS_COLS)
def subfile(work,vid):
    fs=[p for p in work.glob(f"{vid}.*.json3") if "live_chat" not in p.name]
    def score(p):
        n=p.name.lower();return 5 if ".pt-orig." in n else 4 if ".pt-br." in n else 3 if ".pt." in n else 2 if ".en." in n else 1
    return sorted(fs,key=lambda p:(score(p),p.stat().st_mtime),reverse=True)[0] if fs else None
def collect_transcript(url,work,vid,info,browser,attempts,backoff):
    official=info.get("subtitles") or {};auto=info.get("automatic_captions") or {};langs=[x for x in ["pt-orig","pt-BR","pt","pt-PT"] if x in official or x in auto] or ["pt-orig","pt"]
    for lang in langs:
        p=run_retry(ya(browser)+["--skip-download","--no-playlist","--write-subs","--write-auto-subs","--sub-format","json3/best","--sub-langs",lang,"-o",str(work/"%(id)s.%(ext)s"),url],work,f"transcricao {lang}",attempts,backoff)
        f=subfile(work,vid)
        if f and f.stat().st_size:return parse_json3(f),f.name,"rede" if p.returncode==0 else "cache"
    return pd.DataFrame(columns=RAW_TRANS_COLS),"","ausente"
def whisper_fallback(url,work,browser,model,language,attempts,backoff):
    audio=list(work.glob("audio.*"))
    if not audio:
        p=run_retry(ya(browser)+["--no-playlist","-f","bestaudio/best","-x","--audio-format","m4a","-o",str(work/"audio.%(ext)s"),url],work,"audio whisper",attempts,backoff)
        if p.returncode:raise RuntimeError(sanitize_log_detail(p.stdout))
        audio=list(work.glob("audio.*"))
    if not audio:raise RuntimeError("Áudio não localizado")
    from faster_whisper import WhisperModel
    m=WhisperModel(model,device="auto",compute_type="auto");segs,_=m.transcribe(str(audio[0]),language=None if language=="auto" else language,vad_filter=True,beam_size=5)
    return pd.DataFrame([[float(s.start),float(s.end),seconds_to_hms(s.start),seconds_to_hms(s.end),s.text.strip(),f"whisper_{model}"] for s in segs],columns=RAW_TRANS_COLS)
def walk(o):
    if isinstance(o,dict):
        for k,v in o.items():
            if isinstance(v,dict):yield k,v;yield from walk(v)
            elif isinstance(v,list):yield from walk(v)
    elif isinstance(o,list):
        for x in o:yield from walk(x)
def text_obj(x):
    if isinstance(x,str):return x
    if not isinstance(x,dict):return ""
    if "simpleText" in x:return str(x["simpleText"])
    out=[]
    for r in x.get("runs",[]) or []:
        if "text" in r:out.append(str(r["text"]))
        elif "emoji" in r:
            e=r["emoji"];out.append((e.get("shortcuts") or [e.get("emojiId","[emoji]")])[0])
    return "".join(out)
def parse_paid(raw):
    s=str(raw or "").strip();m=re.search(r"([A-Z]{3}|R\$|US\$|€|£|\$)\s*([\d.,]+)",s)
    if not m:return raw,""
    cur={"R$":"BRL","US$":"USD","$":"USD","€":"EUR","£":"GBP"}.get(m.group(1),m.group(1));num=m.group(2)
    if "," in num and "." in num:num=num.replace(".","").replace(",",".")
    elif "," in num:num=num.replace(",",".")
    try:return float(num),cur
    except:return raw,cur
def parse_chat(work,vid):
    fs=list(work.glob(f"{vid}.live_chat*"))+list(work.glob("*live_chat*.json"));path=next((p for p in sorted(fs,key=lambda x:x.stat().st_mtime,reverse=True) if p.is_file() and p.stat().st_size),None)
    if not path:return pd.DataFrame(columns=RAW_CHAT_COLS)
    types={"liveChatTextMessageRenderer":"mensagem","liveChatPaidMessageRenderer":"superchat","liveChatPaidStickerRenderer":"sticker_pago","liveChatMembershipItemRenderer":"membro","liveChatSponsorshipsGiftPurchaseAnnouncementRenderer":"presente_membros","liveChatSponsorshipsGiftRedemptionAnnouncementRenderer":"resgate_presente","liveChatViewerEngagementMessageRenderer":"sistema","liveChatModeChangeMessageRenderer":"sistema"};rows=[]
    for line in path.read_text(encoding="utf-8",errors="replace").splitlines():
        try:o=json.loads(line)
        except:continue
        oms=o.get("videoOffsetTimeMsec") or o.get("replayChatItemAction",{}).get("videoOffsetTimeMsec")
        try:off=float(oms)/1000 if oms is not None else None
        except:off=None
        for k,r in walk(o):
            if k not in types:continue
            msg=text_obj(r.get("message")) or text_obj(r.get("headerSubtext")) or text_obj(r.get("primaryText"))
            if not msg and k=="liveChatPaidStickerRenderer":msg="[sticker pago]"
            amount,currency=parse_paid(text_obj(r.get("purchaseAmountText")))
            rows.append([off,seconds_to_hms(off),r.get("timestampUsec"),types[k],text_obj(r.get("authorName")),r.get("authorExternalChannelId",""),msg,amount,currency,r.get("id","")])
    return pd.DataFrame(rows,columns=RAW_CHAT_COLS).sort_values(["offset_segundos","timestamp_usec"],na_position="last",kind="stable").reset_index(drop=True)
def collect_chat(url,work,vid,browser,attempts,backoff):
    p=run_retry(ya(browser)+["--skip-download","--no-playlist","--write-subs","--sub-langs","live_chat","--sub-format","json/best","-o",str(work/"%(id)s.%(ext)s"),url],work,"chat",attempts,backoff)
    c=parse_chat(work,vid);return c,("rede" if p.returncode==0 else "cache" if not c.empty else "ausente")
def style(path):
    wb=load_workbook(path);fill=PatternFill("solid",fgColor="1F4E78");font=Font(color="FFFFFF",bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
        for c in ws[1]:c.fill=fill;c.font=font;c.alignment=Alignment(horizontal="center",vertical="center")
        for i,col in enumerate(ws.columns,1):
            ml=12
            for j,c in enumerate(col):
                if j>250:break
                ml=max(ml,min(len(str(c.value or ""))+2,55))
                if len(str(c.value or ""))>80:c.alignment=Alignment(vertical="top",wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width=ml
    wb.save(path)
    wb.close()
def sha256_file(path):
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for b in iter(lambda:f.read(1024*1024),b""):h.update(b)
    return h.hexdigest()
def runtime_versions():
    try:
        import yt_dlp;yv=getattr(yt_dlp.version,"__version__","?")
    except:yv="?"
    return {"python":platform.python_version(),"pandas":pd.__version__,"yt_dlp":yv}

def main():
    ap=argparse.ArgumentParser(description="YouTube ETL V7 LIVE FINAL - transcrição + chat ao vivo, sem API key")
    ap.add_argument("url",nargs="?",help="Link do vídeo; se omitido, o programa solicita que você cole o link.")
    ap.add_argument("--saida")
    ap.add_argument("--cookies-from-browser",choices=["chrome","edge","firefox","brave"])
    ap.add_argument("--tentativas",type=int,default=4)
    ap.add_argument("--backoff",type=int,default=15)
    ap.add_argument("--sem-whisper",action="store_true")
    ap.add_argument("--whisper-model",default="small",choices=["tiny","base","small","medium","large-v3"])
    ap.add_argument("--whisper-language",default="pt")
    ap.add_argument("--usar-cache",action="store_true")
    ap.add_argument("--verbose",action="store_true")
    args=ap.parse_args();setup(args.verbose)
    try:
        args.url, requested_vid = normalizar_link(args.url if args.url else solicitar_link())
    except ValueError as exc:
        ap.error(str(exc))

    # Integração HTTP explícita para validar o vídeo antes do ETL principal.
    try:
        dados_http = validar_video_youtube(args.url)
        LOG.info("Integração web OK: %s", dados_http.get("titulo_oembed", "vídeo validado"))
    except Exception as exc:
        LOG.warning("Validação HTTP complementar indisponível: %s", exc)
        dados_http = {}

    vid=requested_vid;cache=Path(__file__).resolve().parent/".youtube_etl_cache"/vid
    if cache.exists() and not args.usar_cache:shutil.rmtree(cache)
    for n in ["metadata","transcript","chat","audio"]:(cache/n).mkdir(parents=True,exist_ok=True)
    logs=[];states={}

    try:
        info=metadata(args.url,cache/"metadata",args.cookies_from_browser,args.tentativas,args.backoff)
        logs.append(["metadados","ok",1,"Metadados obtidos."])
        states["metadados"]={"status":"ok","mandatory":True,"method":"yt-dlp","detail":"Metadados obtidos"}
    except Exception as e:
        info={"id":vid,"title":vid}
        logs.append(["metadados","falhou",0,sanitize_log_detail(e)])
        states["metadados"]={"status":"falhou","mandatory":True,"method":"yt-dlp","detail":sanitize_log_detail(e)}

    vid=str(info.get("id") or vid)
    try:
        tr,source,origin=collect_transcript(args.url,cache/"transcript",vid,info,args.cookies_from_browser,args.tentativas,args.backoff)
        logs.append(["transcricao","ok" if len(tr) else "vazia",len(tr),f"Fonte={source}; origem={origin}"])
        states["transcricao"]={"status":"ok" if len(tr) else "falhou","mandatory":True,"method":"yt-dlp_legendas","detail":f"Fonte={source}; origem={origin}"}
    except Exception as e:
        tr=pd.DataFrame(columns=RAW_TRANS_COLS)
        logs.append(["transcricao","falhou",0,sanitize_log_detail(e)])
        states["transcricao"]={"status":"falhou","mandatory":True,"method":"yt-dlp_legendas","detail":sanitize_log_detail(e)}

    if tr.empty and not args.sem_whisper:
        try:
            tr=whisper_fallback(args.url,cache/"audio",args.cookies_from_browser,args.whisper_model,args.whisper_language,args.tentativas,args.backoff)
            logs.append(["transcricao_whisper","ok",len(tr),f"Modelo={args.whisper_model}"])
            states["transcricao"]={"status":"ok","mandatory":True,"method":"faster-whisper","detail":f"Modelo={args.whisper_model}"}
        except Exception as e:
            logs.append(["transcricao_whisper","falhou",0,sanitize_log_detail(e)])

    try:
        cr,corigin=collect_chat(args.url,cache/"chat",vid,args.cookies_from_browser,args.tentativas,args.backoff)
        is_live_recording=bool(info.get("was_live") or info.get("live_status") in ("was_live","is_live"))
        chat_status="ok" if len(cr) else ("falhou" if is_live_recording else "indeterminado")
        logs.append(["chat",chat_status,len(cr),f"Replay processado; origem={corigin}"])
        states["chat"]={"status":chat_status,"mandatory":is_live_recording,"method":"yt-dlp_live_chat","detail":f"origem={corigin}"}
    except Exception as e:
        cr=pd.DataFrame(columns=RAW_CHAT_COLS)
        is_live_recording=bool(info.get("was_live") or info.get("live_status") in ("was_live","is_live"))
        logs.append(["chat","falhou",0,sanitize_log_detail(e)])
        states["chat"]={"status":"falhou","mandatory":is_live_recording,"method":"yt-dlp_live_chat","detail":sanitize_log_detail(e)}

    ct,stats=clean_chat(cr,info);tt=treat_transcript(tr,vid);timeline=build_timeline(tt,ct,vid)
    q,status=build_quality(info,tr,tt,cr,ct,stats,timeline,states)
    versions=runtime_versions()
    md=pd.DataFrame([["pipeline_version",PIPELINE_VERSION],["url",args.url],["video_id",vid],["titulo",info.get("title","")],["canal",info.get("channel","") or info.get("uploader","")],["canal_id",info.get("channel_id","")],["data_upload",info.get("upload_date","")],["duracao_segundos",info.get("duration","")],["visualizacoes",info.get("view_count","")],["curtidas",info.get("like_count","")],["foi_ao_vivo",info.get("was_live","")],["status_live",info.get("live_status","")],["idioma",info.get("language","")],["python_version",versions["python"]],["pandas_version",versions["pandas"]],["yt_dlp_version",versions["yt_dlp"]],["cache_reutilizado",args.usar_cache],["data_hora_coleta",datetime.now(timezone.utc).isoformat(timespec="seconds")],
        ["titulo_validacao_http",dados_http.get("titulo_oembed","")],
        ["autor_validacao_http",dados_http.get("autor_oembed","")]],columns=["campo","valor"])
    resumo=pd.DataFrame([["Status de produção",status],["Vídeo",info.get("title",vid)],["Duração (s)",info.get("duration","")],["Segmentos de transcrição",len(tr)],["Blocos tratados",len(tt)],["Chat bruto",len(cr)],["Chat tratado",len(ct)],["Participantes do público identificados",stats.get("participantes_externos_unicos",stats.get("participantes_unicos",0))],["Mensagens durante o evento",stats.get("mensagens_durante_evento",0)],["Offset do chat confiável","SIM" if stats.get("offset_confiavel") else "NAO"]],columns=["Métrica","Valor"])

    manifest=[]
    for stage in ["transcript","chat"]:
        for f in sorted((cache/stage).glob("*")):
            if f.is_file():manifest.append([stage,f.name,f.stat().st_size,sha256_file(f)])
    manifest=pd.DataFrame(manifest,columns=["etapa","arquivo_fonte","bytes","sha256"]);fontes=source_status_table(states);now=datetime.now(timezone.utc).isoformat(timespec="seconds")
    logdf=pd.DataFrame([[PIPELINE_VERSION,now,*x] for x in logs]+[[PIPELINE_VERSION,now,"gate_producao","ok" if status=="APROVADO_PRODUCAO" else "falhou",1,status]],columns=["pipeline_version","data_hora","etapa","status","registros","detalhe"])

    # Informações adicionais sem remover os campos do contrato V7.
    extra = [["release_timestamp", info.get("release_timestamp", "")],
             ["live_start_time", info.get("live_start_time", "")],
             ["live_end_time", info.get("live_end_time", "")],
             ["descricao", info.get("description", "")],
             ["relatorio_version", "1.0.0"]]
    md = pd.concat([md, pd.DataFrame(extra, columns=["campo", "valor"])], ignore_index=True)
    # Preserva também o JSON completo de metadados para auditoria futura.
    metadata_file = cache/"metadata"/f"{vid}.metadata.json"
    metadata_file.write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest.loc[len(manifest)] = ["metadata", metadata_file.name, metadata_file.stat().st_size, sha256_file(metadata_file)]
    out=Path(args.saida).resolve() if args.saida else Path.cwd()/f"YouTube_{vid}_GERENCIAL.xlsx"
    tmp=criar_temporario(out)
    with pd.ExcelWriter(tmp,engine="openpyxl") as w:
        resumo.to_excel(w,index=False,sheet_name="Resumo")
        md.to_excel(w,index=False,sheet_name="Metadados")
        tr.to_excel(w,index=False,sheet_name="Transcricao")
        tt.to_excel(w,index=False,sheet_name="Transcricao_tratada")
        cr.to_excel(w,index=False,sheet_name="Chat_ao_vivo")
        ct.to_excel(w,index=False,sheet_name="Chat_tratado")
        timeline.to_excel(w,index=False,sheet_name="Linha_do_tempo")
        q.to_excel(w,index=False,sheet_name="Qualidade_Coleta")
        fontes.to_excel(w,index=False,sheet_name="Status_Fontes")
        manifest.to_excel(w,index=False,sheet_name="Manifesto_Fontes")
        logdf.to_excel(w,index=False,sheet_name="Log_coleta")
    # No contrato atual as abas técnicas são valores, não fórmulas.
    # Impede que mensagens iniciadas por '=' sejam executadas pelo Excel.
    safe_book = load_workbook(tmp)
    for sheet in safe_book:
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    safe_book.save(tmp);safe_book.close()
    style(tmp)
    try:
        adicionar_relatorio(tmp)
    except Exception:
        # O Excel técnico já produzido fica recuperável mesmo se o painel falhar.
        fallback = out.with_name(out.stem + "_BASE_TECNICA.xlsx")
        fallback = publicar_excel(tmp, fallback)
        LOG.exception("Falha na apresentação. Base técnica preservada em %s", fallback)
        return 1
    out=publicar_excel(tmp,out);print(f"Excel salvo em: {out}");print(status);return 0 if status=="APROVADO_PRODUCAO" else 3
if __name__=="__main__":raise SystemExit(main())
