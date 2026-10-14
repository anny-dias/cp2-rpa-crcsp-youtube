"""Camada de apresentação. Não altera coleta, tratamento ou abas de origem.

Os agregados são calculados em Python e publicados como um retrato da coleta.
Para atualizar o relatório, execute novamente gerar_relatorio.py.
"""
from __future__ import annotations

import math
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from etl_core import canonical_author_type, participant_identity_metrics

REPORT_VERSION = '1.1.0'
SHEETS = ['Painel', 'Indicadores', 'Analise_chat', 'Picos_5_min',
          'Mensagens_analise', 'Participantes_chat', 'Leitura_transcricao',
          'Validacao_relatorio', 'Metodologia']
NAVY, BLUE, LIGHT, AMBER = '17365D', '2F75B5', 'EAF1F8', 'FFF2CC'
NA = 'Não disponível'

# Categorias mutuamente exclusivas, na ordem de prioridade documentada.
# Pergunta é um sinal separado: pode coexistir com qualquer categoria.
RULES = [
    ('Áudio/conexão/imagem', r'\b(audio|som|microfone|conexao|travando|travou|sem imagem|nao (?:consigo )?(?:ouvir|escutar))\b', True),
    ('Certificado/EPC', r'\b(certificado\w*|epc|pontuacao|pontos)\b', True),
    ('Material/gravação', r'\b(material|materiais|slides?|apostila\w*|gravacao|gravado|apresentacao|pdf)\b', True),
    ('Acesso/presença/inscrição', r'\b(presenca|inscri\w*|acesso|lista|link|comec\w*|inicio)\b', True),
    ('Contato do palestrante', r'\b(contato|email|e-mail|instagram|telefone|whatsapp)\b', False),
    ('Crítica/sugestão', r'\b(sugiro|sugestao|sugestoes|critica|ruim|pessimo\w*)\b', False),
    ('Elogio/agradecimento', r'\b(obrigad\w*|agradec\w*|parabens|excelente\w*|otim\w*|maravilh\w*)\b', False),
    ('Cumprimento', r'\b(bom dia|boa tarde|boa noite|ola|oi)\b', False),
]
CATEGORIES = [r[0] for r in RULES] + ['Outro/indeterminado']
QUESTION = r'(?:\?|\b(?:como|qual|quais|quando|onde|por que|porque|quanto\w*|quem|duvida\w*)\b|\b(?:e possivel|tem como|gostaria de saber|queria saber|quero saber|alguem sabe|saber se|sera que|ha possibilidade|haveria possibilidade)\b|\b(?:posso|podemos|devo|devemos)\b|\b(?:pode|podem|deve|devem)\s+(?:ser|usar|utilizar|fazer|considerar|aplicar|informar|emitir|lancar|calcular|deduzir|compensar|registrar|entregar|enviar|adotar|reconhecer|contabilizar)\b)'
PERIODS = ['Pré-evento', 'Q1', 'Q2', 'Q3', 'Q4', 'Pós-evento', 'Indeterminado']


def text(value):
    return '' if value is None else str(value).strip()


def number(value):
    try:
        n = float(value)
        return n if math.isfinite(n) else None
    except (TypeError, ValueError):
        return None


def nonnegative(value):
    n = number(value)
    return n if n is not None and n >= 0 else None


def display_date(value, timestamp=False):
    try:
        if timestamp:
            return datetime.fromisoformat(text(value)).strftime('%d/%m/%Y %H:%M %z')
        return datetime.strptime(text(value), '%Y%m%d').strftime('%d/%m/%Y')
    except ValueError:
        return text(value) or NA


def norm(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', text(value).lower())
                   if not unicodedata.combining(c))


def records(wb, name):
    if name not in wb:
        return []
    rows = iter(wb[name].values)
    headers = next(rows, ())
    return [dict(zip(headers, row), _linha=i) for i, row in enumerate(rows, 2)
            if any(v is not None for v in row)]


def classify(message):
    value = norm(message)
    question = bool(re.search(QUESTION, value))
    for category, pattern, operational in RULES:
        if re.search(pattern, value):
            return category, question, operational
    return 'Outro/indeterminado', question, False


def period(offset, duration, reliable):
    if offset is None or not reliable:
        return 'Indeterminado'
    if offset < 0:
        return 'Pré-evento'
    if not duration or duration <= 0:
        return 'Indeterminado'
    if offset > duration:
        return 'Pós-evento'
    return 'Q' + str(min(4, int(offset * 4 / duration) + 1))


def analyze(wb):
    required = {'Metadados', 'Chat_ao_vivo', 'Chat_tratado', 'Transcricao', 'Transcricao_tratada'}
    missing = required.difference(wb.sheetnames)
    if missing:
        raise ValueError('Abas obrigatórias ausentes: ' + ', '.join(sorted(missing)))
    md = {r.get('campo'): r.get('valor') for r in records(wb, 'Metadados')}
    quality = {r.get('indicador'): r.get('valor') for r in records(wb, 'Qualidade_Coleta')}
    duration = number(md.get('duracao_segundos'))
    duration = duration if duration is not None and duration > 0 else None
    reliable = text(quality.get('chat_offset_confiavel')).upper() == 'SIM'
    cid = text(md.get('canal_id'))
    raw, treated = records(wb, 'Chat_ao_vivo'), records(wb, 'Chat_tratado')
    raw_by_id = defaultdict(list)
    raw_by_tuple = defaultdict(list)
    for r in raw:
        if text(r.get('id_mensagem')):
            raw_by_id[text(r['id_mensagem'])].append(r['_linha'])
        raw_by_tuple[tuple(text(r.get(c)) for c in ['timestamp_usec', 'autor', 'mensagem', 'tipo'])].append(r['_linha'])
    analysis = []
    for r in treated:
        aid = text(r.get('canal_autor_id'))
        atype = canonical_author_type(r.get('autor'), aid, cid)
        population = ('Sistema' if text(r.get('tipo')) == 'sistema' else
                      'Canal oficial' if atype == 'canal_oficial' else
                      'Equipe/conta institucional' if atype == 'equipe_crcsp' else
                      'Público externo' if aid else 'Público externo sem ID')
        category, question, operational = classify(r.get('mensagem'))
        offset = number(r.get('offset_segundos'))
        origin = raw_by_id.get(text(r.get('id_mensagem'))) or raw_by_tuple.get(
            tuple(text(r.get(c)) for c in ['timestamp_usec', 'autor', 'mensagem', 'tipo']), [])
        analysis.append({**r, 'populacao': population, 'categoria': category,
                         'pergunta': question, 'operacional': operational,
                         'offset': offset, 'faixa': period(offset, duration, reliable),
                         'linhas_brutas': ', '.join(map(str, origin))})
    public = [r for r in analysis if r['populacao'] in {'Público externo', 'Público externo sem ID'}]
    official = [r for r in analysis if r['populacao'] == 'Canal oficial']
    institutional = [r for r in analysis if r['populacao'] == 'Equipe/conta institucional']
    participants = defaultdict(list)
    for r in public:
        if text(r.get('canal_autor_id')):
            participants[text(r['canal_autor_id'])].append(r)
    identity = participant_identity_metrics(pd.DataFrame(analysis), cid)
    # Defesa contra regressão: a mesma função canônica deve reconciliar com o agrupamento gerencial.
    if identity['participantes_publico_ids'] != len(participants):
        raise ValueError('Divergência interna na métrica de participantes do público.')
    counts = Counter(r['categoria'] for r in public)
    phases = Counter(r['faixa'] for r in public)
    bins = defaultdict(lambda: [0, 0, 0, 0])
    for r in analysis:
        if r['offset'] is None or not reliable or r['populacao'] == 'Sistema':
            continue
        b = math.floor(r['offset'] / 300) * 5
        v = bins[b]
        if r['populacao'] in {'Canal oficial', 'Equipe/conta institucional'}:
            v[3] += 1
        elif r['populacao'] in {'Público externo', 'Público externo sem ID'}:
            v[0] += 1
            v[1] += int(r['pergunta'])
            v[2] += int(r['operacional'])
    # Inclui intervalos sem mensagens para não comprimir silenciosamente o tempo.
    if reliable and duration and duration <= 31 * 86400:
        first = min([0, *bins.keys()])
        last = max([math.floor(duration / 300) * 5, *bins.keys()])
        if last - first <= 31 * 24 * 60:
            for b in range(first, last + 5, 5):
                bins[b]
    n = len(public)
    hourly_n = sum(r['faixa'] in {'Q1', 'Q2', 'Q3', 'Q4'} for r in public)
    source_status = {norm(r.get('fonte')): text(r.get('status')).upper()
                     for r in records(wb, 'Status_Fontes')}
    chat_available = bool(analysis) or source_status.get('chat ao vivo') == 'OK'
    indicators = [
        ('Visualizações acumuladas', nonnegative(md.get('visualizacoes')), 'Metadados.visualizacoes; acumuladas até a coleta', '#,##0'),
        ('Curtidas acumuladas', nonnegative(md.get('curtidas')), 'Metadados.curtidas; acumuladas até a coleta', '#,##0'),
        ('Duração (horas)', duration / 3600 if duration else None, 'Metadados.duracao_segundos / 3600', '0.00'),
        ('Mensagens do público externo', n if chat_available else None, 'Chat_tratado excluindo canal oficial, contas institucionais/equipe e sistema; inclui registros públicos sem ID', '#,##0'),
        ('Mensagens do canal oficial', len(official) if chat_available and cid else None, 'ID do autor igual a Metadados.canal_id', '#,##0'),
        ('Mensagens de contas institucionais/equipe', len(institutional) if chat_available else None, 'Contas com nome público associado ao CRCSP, separadas do público', '#,##0'),
        ('Participantes identificáveis', identity['participantes_publico_ids'] if chat_available else None, 'Métrica canônica por canal_autor_id; exclui canal oficial, contas institucionais/equipe e registros sem ID', '#,##0'),
        ('Perguntas sinalizadas automaticamente', sum(r['pergunta'] for r in public) if chat_available else None, "Sinal ? ou padrão linguístico de pergunta; não validada", '#,##0'),
        ('Menções operacionais automáticas', sum(r['operacional'] for r in public) if chat_available else None, 'Categorias áudio, certificado, material e acesso; não comprova demanda', '#,##0'),
        ('Elogios/agradecimentos automáticos', counts['Elogio/agradecimento'] if chat_available else None, 'Categoria exclusiva; não mede satisfação geral', '#,##0'),
        ('Mensagens sem ID de autor', sum(r['populacao'] == 'Público externo sem ID' for r in public) if chat_available else None, 'Incluídas nas mensagens; excluídas de participantes únicos', '#,##0'),
        ('Mensagens por hora durante o vídeo', hourly_n / (duration / 3600) if duration and reliable and chat_available else None, 'Mensagens Q1–Q4 / duração em horas; exclui pré/pós e tempo indeterminado', '0.0'),
        ('Proporção de perguntas sinalizadas', sum(r['pergunta'] for r in public) / n if n else None, 'Perguntas sinalizadas / mensagens do público externo', '0.0%'),
        ('Proporção de menções operacionais', sum(r['operacional'] for r in public) / n if n else None, 'Menções operacionais / mensagens do público externo', '0.0%'),
        ('Proporção de curtidas por visualização', nonnegative(md.get('curtidas')) / nonnegative(md.get('visualizacoes')) if nonnegative(md.get('visualizacoes')) and nonnegative(md.get('curtidas')) is not None else None, 'Curtidas acumuladas / visualizações acumuladas; não mede satisfação', '0.0%'),
        ('Concentração dos 10 IDs mais ativos', sum(sorted([len(v) for v in participants.values()], reverse=True)[:10]) / n if n else None, 'Mensagens dos 10 IDs públicos mais ativos / mensagens do público externo; registros sem ID ficam apenas no denominador', '0.0%'),
        ('Mensagens com tempo indeterminado', phases['Indeterminado'] if chat_available else None, 'Tempo ausente, duração ausente ou diagnóstico temporal não confiável', '#,##0'),
    ]
    return dict(md=md, quality=quality, duration=duration, reliable=reliable, raw=raw,
                analysis=analysis, public=public, official=official, institutional=institutional, participants=participants, identity=identity,
                counts=counts, phases=phases, bins=dict(sorted(bins.items())),
                indicators=indicators, chat_available=chat_available)


def put(ws, row, col, value):
    cell = ws.cell(row, col, value if value is not None else NA)
    # Conteúdo do YouTube é texto, nunca uma fórmula executável.
    if isinstance(value, str):
        cell.data_type = 's'
    return cell


def table(ws, headers, rows, widths, start=4):
    for j, h in enumerate(headers, 1):
        c = put(ws, start, j, h)
        c.font = Font(name='Arial', size=10, color='FFFFFF', bold=True)
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[start].height = 32
    for i, row in enumerate(rows, start + 1):
        max_lines = 1
        for j, value in enumerate(row, 1):
            c = put(ws, i, j, value)
            c.font = Font(name='Arial', size=10, color='243746')
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if i % 2:
                c.fill = PatternFill('solid', fgColor='F2F5F9')
            max_lines = max(max_lines, math.ceil(len(str(value or '')) / max(8, widths[j - 1] - 3)))
        ws.row_dimensions[i].height = min(400, max(22, max_lines * 14 + 6))
    for j, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = f'A{start + 1}'
    if rows:
        ref = f'A{start}:{get_column_letter(len(headers))}{start + len(rows)}'
        used_names = {name for sheet in ws.parent for name in sheet.tables}
        base = 'Gerencial_' + re.sub('[^a-zA-Z0-9]', '', unicodedata.normalize('NFKD', ws.title))
        name = base
        serial = 2
        while name in used_names:
            name = base + str(serial)
            serial += 1
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(t)


def heading(ws, title, note='', end='H'):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = BLUE
    ws.merge_cells(f'A1:{end}1')
    c = put(ws, 1, 1, title)
    c.font = Font(name='Arial', size=18, bold=True, color=NAVY)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f'A2:{end}2')
    c = put(ws, 2, 1, note)
    c.font = Font(name='Arial', size=10, color='536778')
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[2].height = 35
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:4'


def chart(ws, source, title, min_col, max_col, min_row, max_row, category_col,
          anchor, line=False):
    if max_row <= min_row:
        return
    obj = LineChart() if line else BarChart()
    obj.title = title
    obj.style = 13
    obj.add_data(Reference(source, min_col=min_col, max_col=max_col,
                           min_row=min_row, max_row=max_row), titles_from_data=True)
    obj.set_categories(Reference(source, min_col=category_col, min_row=min_row+1, max_row=max_row))
    obj.width, obj.height = (14, 11) if ws.title == 'Painel' else (18, 9)
    obj.y_axis.title = 'Mensagens'
    obj.x_axis.title = 'Início da janela (min)' if line else 'Categoria / período'
    if not line:
        obj.type = 'bar'
        obj.x_axis.title = 'Mensagens'
        obj.y_axis.title = 'Categoria / período'
        obj.dataLabels = DataLabelList()
        obj.dataLabels.showVal = True
    for series in obj.series:
        series.graphicalProperties.solidFill = BLUE
        series.graphicalProperties.line.solidFill = BLUE
        if line:
            series.smooth = False
    if min_col == max_col:
        obj.legend = None
    ws.add_chart(obj, anchor)


def build_report(wb, source_name='Coleta atual'):
    """Adiciona/recria apenas abas geradas por esta camada; valida conflitos."""
    existing = set(SHEETS).intersection(wb.sheetnames)
    if existing and ('Metodologia' not in wb or wb['Metodologia']['A3'].value != 'youtube_gerencial:' + REPORT_VERSION):
        raise ValueError('Conflito com abas existentes: ' + ', '.join(sorted(existing)))
    a = analyze(wb)
    for name in existing:
        del wb[name]
    presentation_names = {
        'Dashboard', 'Resumo por vídeo', 'Resumo da transmissão', 'Participação',
        'Picos de 5 minutos', 'Perguntas de conteúdo', 'Transcrição dos Chats',
        'Resumo_Categorias', 'Ocorrências e sugestões', 'Participantes Públicos',
        'Metodologia e lacunas'
    }
    source_sheets = [name for name in wb.sheetnames if name not in presentation_names]
    generated = {name: wb.create_sheet(name, i) for i, name in enumerate(SHEETS)}
    md, q, public = a['md'], a['quality'], a['public']
    title = text(md.get('titulo')) or text(md.get('video_id')) or 'Vídeo sem título'
    vid = text(md.get('video_id'))
    url = f'https://www.youtube.com/watch?v={vid}' if re.fullmatch(r'[A-Za-z0-9_-]{11}', vid) else ''

    ws = generated['Indicadores']
    heading(ws, 'Indicadores da transmissão', 'Valores calculados na geração do relatório. Critérios e denominadores explícitos.', 'C')
    table(ws, ['Indicador', 'Valor', 'Cálculo / fonte'], [r[:3] for r in a['indicators']], [47, 24, 104])
    for i, row in enumerate(a['indicators'], 5):
        ws.cell(i, 2).number_format = row[3]

    ws = generated['Mensagens_analise']
    heading(ws, 'Mensagens e sinais automáticos', 'Filtros para consulta e triagem. As categorias não foram validadas por uma pessoa.', 'H')
    headers = ['Tempo no vídeo', 'Autor', 'Mensagem', 'População por ID', 'Categoria automática',
               'Pergunta sinalizada', 'Menção operacional', 'Período gerencial', 'ID do autor',
               'Chave do registro', 'Linha Chat_tratado', 'Linhas Chat_ao_vivo', 'Categoria de autor V7', 'Link do vídeo']
    rows = [[r.get('tempo_video'), r.get('autor'), r.get('mensagem'), r['populacao'], r['categoria'],
             'Sim' if r['pergunta'] else 'Não', 'Sim' if r['operacional'] else 'Não', r['faixa'],
             r.get('canal_autor_id'), r.get('chave_evento'), r['_linha'], r['linhas_brutas'],
             r.get('autor_tipo'), url + f'&t={max(0, int(r["offset"]))}s' if url and r['offset'] is not None and a['reliable'] else url]
            for r in a['analysis']]
    table(ws, headers, rows, [19, 28, 95, 23, 31, 21, 22, 22, 30, 38, 21, 28, 24, 52])
    for i, r in enumerate(a['analysis'], 5):
        ws.cell(i, 11).hyperlink = f"#'Chat_tratado'!A{r['_linha']}"
        ws.cell(i, 11).style = 'Hyperlink'
        if ws.cell(i, 14).value:
            ws.cell(i, 14).hyperlink = ws.cell(i, 14).value

    ws = generated['Participantes_chat']
    heading(ws, 'Participantes identificáveis do chat', 'Público externo agrupado por ID. Pessoas sem ID não são deduplicadas pelo nome. Chat não equivale à audiência.', 'F')
    rows = []
    for aid, msgs in sorted(a['participants'].items(), key=lambda item: (-len(item[1]), item[0])):
        offsets = [r['offset'] for r in msgs if r['offset'] is not None and a['reliable']]
        rows.append([aid, ' / '.join(sorted({text(r.get('autor')) for r in msgs})), len(msgs),
                     sum(r['pergunta'] for r in msgs), min(offsets) / 60 if offsets else None,
                     max(offsets) / 60 if offsets else None])
    table(ws, ['ID público', 'Nome(s) na fonte', 'Mensagens', 'Perguntas sinalizadas', 'Primeira mensagem (min)', 'Última mensagem (min)'], rows, [35, 44, 18, 25, 25, 25])
    for row in ws.iter_rows(min_row=5, min_col=5, max_col=6):
        for c in row:
            c.number_format = '0.0'

    ws = generated['Analise_chat']
    heading(ws, 'Natureza e distribuição do chat', 'Público externo, incluindo registros sem ID. Sinais automáticos sujeitos a revisão. Dados ausentes não são zeros.', 'H')
    rows = [[c, a['counts'][c] if a['chat_available'] else None,
             a['counts'][c] / len(public) if public else None] for c in CATEGORIES]
    table(ws, ['Categoria automática exclusiva', 'Mensagens', '% do público externo'], rows, [35, 20, 24])
    for i in range(5, 5 + len(rows)):
        ws.cell(i, 3).number_format = '0.0%'
    phase_start = 18
    rows = [[p, a['phases'][p] if a['chat_available'] else None] for p in PERIODS]
    table(ws, ['Período do vídeo', 'Mensagens'], rows, [35, 20], phase_start)
    ws.freeze_panes = 'A5'
    put(ws, 28, 1, 'Fonte: Mensagens_analise. Categorias exclusivas; pergunta é um sinal independente.')
    ws.merge_cells('A28:L29')
    ws['A28'].alignment = Alignment(wrap_text=True)
    if a['chat_available']:
        chart(ws, ws, 'Categorias automáticas', 2, 2, 4, 4+len(CATEGORIES), 1, 'E4')
        chart(ws, ws, 'Distribuição temporal do chat', 2, 2, phase_start, phase_start+len(PERIODS), 1, 'E22')

    ws = generated['Picos_5_min']
    heading(ws, 'Dinâmica em janelas de 5 minutos', 'Intervalos fixos [início, fim), alinhados ao início do vídeo. Pré/pós incluídos. Os picos não medem retenção.', 'G')
    rows = [[b, b+5, *values] for b, values in a['bins'].items()]
    table(ws, ['Início (min)', 'Fim exclusivo (min)', 'Público externo', 'Perguntas sinalizadas', 'Menções operacionais', 'Canal oficial'], rows, [18, 23, 22, 26, 26, 22])
    if not a['reliable']:
        put(ws, 5, 1, 'Indisponível: tempo do chat não confirmado pela qualidade da coleta.')
    elif rows and a['chat_available']:
        chart(ws, ws, 'Mensagens ao longo do vídeo', 3, 3, 4, 4+len(rows), 1, 'H4', True)

    ws = generated['Leitura_transcricao']
    heading(ws, 'Transcrição para leitura', 'Blocos do tratamento original, com ligação à fonte. Transcrição automática pode conter erros e não comprova cobertura do programa.', 'F')
    trans = records(wb, 'Transcricao_tratada')
    rows = [[r.get('inicio'), r.get('fim'), r.get('texto'), r.get('fonte'), r.get('chave_bloco'), r['_linha']] for r in trans]
    table(ws, ['Início', 'Fim', 'Texto', 'Fonte', 'Chave do bloco', 'Linha de origem'], rows, [19, 19, 120, 24, 38, 20])
    for i, r in enumerate(trans, 5):
        ws.cell(i, 6).hyperlink = f"#'Transcricao_tratada'!A{r['_linha']}"
        ws.cell(i, 6).style = 'Hyperlink'

    # Conciliação independente dos totais de apresentação.
    ws = generated['Validacao_relatorio']
    heading(ws, 'Conciliação e rastreabilidade', 'Checks do relatório não substituem o gate original de qualidade da coleta.', 'E')
    unique_raw, seen, dup = [], set(), 0
    for r in a['raw']:
        mid = text(r.get('id_mensagem'))
        key = ('id', mid) if mid else ('fallback', *(text(r.get(c)) for c in ['timestamp_usec', 'autor', 'mensagem', 'tipo']))
        if key in seen:
            dup += 1
        else:
            seen.add(key)
            unique_raw.append(r)
    system = sum(text(r.get('tipo')) == 'sistema' for r in unique_raw)
    checks = [
        ('Bruto = duplicidades + sistema + tratado', len(a['raw']), dup + system + len(a['analysis']), 'Chat_ao_vivo / Chat_tratado'),
        ('Tratado = populações gerenciais', len(a['analysis']), sum(Counter(r['populacao'] for r in a['analysis']).values()), 'População calculada por ID'),
        ('Categorias = público externo', len(public), sum(a['counts'].values()), 'Categorias mutuamente exclusivas'),
        ('Períodos = público externo', len(public), sum(a['phases'].values()), 'Inclui tempo indeterminado'),
        ('Mensagens por ID + sem ID = público externo', len(public), sum(len(v) for v in a['participants'].values()) + sum(r['populacao'] == 'Público externo sem ID' for r in public), 'Contagem de mensagens, não de pessoas'),
        ('Participantes = IDs únicos do público externo', a['identity']['participantes_publico_ids'], len({text(r.get('canal_autor_id')) for r in public if text(r.get('canal_autor_id'))}), 'Métrica canônica participant_identity_metrics'),
        ('Populações institucionais separadas do público', len(a['official']) + len(a['institutional']), sum(r['populacao'] in {'Canal oficial','Equipe/conta institucional'} for r in a['analysis']), 'Conciliação por mensagens'),
        ('IDs identificados = público + institucionais', a['identity']['autores_identificados_total'], a['identity']['participantes_publico_ids'] + a['identity']['institucionais_ids'], 'Conciliação por IDs únicos'),
        ('Registros com origem no chat bruto', len(a['analysis']), sum(bool(r['linhas_brutas']) for r in a['analysis']), 'ID da mensagem; fallback timestamp/autor/mensagem/tipo'),
        ('Chaves tratadas únicas e preenchidas', len(a['analysis']), len({text(r.get('chave_evento')) for r in a['analysis'] if text(r.get('chave_evento'))}), 'Chat_tratado.chave_evento'),
    ]
    if a['reliable']:
        checks.append(('Janelas = público externo com tempo', sum(r['offset'] is not None for r in public), sum(v[0] for v in a['bins'].values()), 'Intervalos fixos de 300 segundos'))
    checks.append(('Metadados numéricos negativos', 0, sum(number(md.get(k)) is not None and number(md.get(k)) < 0 for k in ['visualizacoes', 'curtidas', 'duracao_segundos']), 'Valores negativos preservados na fonte e indisponíveis nos indicadores'))
    table(ws, ['Conciliação', 'Esperado', 'Calculado', 'Resultado', 'Critério'],
          [[label, expected, actual, 'OK' if expected == actual else 'DIVERGÊNCIA', detail]
           for label, expected, actual, detail in checks], [55, 18, 18, 22, 85])
    for i in range(5, 5+len(checks)):
        if ws.cell(i, 4).value != 'OK':
            ws.cell(i, 4).fill = PatternFill('solid', fgColor='FCE4D6')
    validation_ok = all(x[1] == x[2] for x in checks)

    ws = generated['Metodologia']
    heading(ws, 'Método, fontes e limitações', 'Relatório de um vídeo por execução. Não há avaliação automática de qualidade do evento.', 'D')
    put(ws, 3, 1, 'youtube_gerencial:' + REPORT_VERSION)
    method = [
        ('Arquivo de origem', source_name, 'Todas as abas originais permanecem no arquivo.', 'Fonte'),
        ('Gerado em UTC', datetime.now(timezone.utc).isoformat(timespec='seconds'), 'A data de geração não substitui a data da coleta.', 'Processamento'),
        ('Data da coleta', text(md.get('data_hora_coleta')) or NA, 'Visualizações e curtidas são acumuladas nessa data.', 'Fonte'),
        ('Cálculos', 'Retrato calculado em Python; fórmulas descritas em Indicadores.', 'Regenere o relatório após alterar os dados. Não há atualização automática dos agregados no Excel.', 'Método'),
        ('População', 'Canal oficial = ID exato do canal; contas com nome público CRCSP ficam em grupo institucional/equipe; público externo = autores restantes após excluir esses grupos.', 'A identificação da equipe por nome público é heurística e auditável; registros sem ID entram em mensagens, mas não em participantes únicos.', 'Método'),
        ('Identidade', 'Participantes = IDs distintos; nomes não são chave.', 'Um ID não comprova identidade civil. Não equivale a espectadores.', 'Método'),
        ('Classificação', 'Primeira regra correspondente define uma categoria exclusiva. Sem correspondência: Outro/indeterminado.', 'Menções são sinais exploratórios, sujeitos a falsos positivos, negações e contexto.', 'Automático'),
        ('Perguntas sinalizadas', 'Sinal ? ou padrões linguísticos de intenção interrogativa: ' + QUESTION, 'Pode coexistir com qualquer categoria. Não usar como quantidade oficial de perguntas técnicas nem taxa de resposta.', 'Automático'),
        ('Menções operacionais', 'Categorias áudio/conexão, certificado/EPC, material/gravação, acesso/presença/inscrição.', 'Menção não comprova dificuldade, solicitação ou incidente.', 'Automático'),
        ('Períodos', 'Pré: t<0; Q1: [0,D/4); Q2: [D/4,D/2); Q3: [D/2,3D/4); Q4: [3D/4,D]; pós: t>D.', 'Exige duração positiva e chat_offset_confiavel=SIM. Demais casos: indeterminado.', 'Método'),
        ('Picos', 'Janelas fixas de 300 s com floor(t/300), inclusive offsets negativos.', 'Não é janela móvel. Empates: primeira janela cronológica. Não mede retenção.', 'Método'),
        ('Bases pequenas', 'Exibe sempre os números absolutos junto aos percentuais; alerta quando N<30.', 'Corte descritivo de conveniência, não teste estatístico nem nota de qualidade.', 'Limitação'),
        ('Dados indisponíveis', 'Retenção, espectadores únicos, satisfação, aprendizagem, presença e certificados emitidos.', 'Chat, curtidas e visualizações não substituem essas dimensões.', 'Lacuna'),
        ('Escopo individual', 'Um vídeo por execução.', 'Todos os indicadores deste arquivo se referem exclusivamente à transmissão selecionada.', 'Escopo'),
        ('Data do vídeo', 'data_upload é publicação; não é necessariamente início efetivo da live.', 'Início real só aparece quando o metadado correspondente existe.', 'Limitação'),
        ('Transmissão em andamento', 'Os resultados representam apenas os dados disponíveis na coleta.', 'Uma live em andamento ou agendada não deve ser apresentada como varredura final completa.', 'Limitação'),
        ('Status V7', text(q.get('status_geral')) or NA, 'Aprovação do gate técnico não valida classificações gerenciais nem comprova integralidade absoluta.', 'Coleta'),
    ]
    method.extend((f'Regra {i}: {cat}', regex, 'Operacional' if op else 'Não operacional', 'Automático')
                  for i, (cat, regex, op) in enumerate(RULES, 1))
    method.extend((f'Aba original: {name}', f'{wb[name].max_row - 1} linhas de dados; {wb[name].max_column} colunas',
                   'Preservada com o nome e valores de origem.', 'Fonte') for name in source_sheets)
    table(ws, ['Item', 'Definição / regra / fonte', 'Limitação / utilização', 'Natureza'], method, [35, 100, 96, 20])

    ws = generated['Painel']
    heading(ws, 'Análise da transmissão', title, 'L')
    ws.sheet_properties.tabColor = NAVY
    for col in range(1, 13):
        ws.column_dimensions[get_column_letter(col)].width = 13
    for row in range(3, 65):
        ws.row_dimensions[row].height = 22
    gate = text(q.get('status_geral')) or NA
    status = text(md.get('status_live'))
    gate_label = {'APROVADO_PRODUCAO': 'aprovada nos controles técnicos', 'NAO_APROVADO_PRODUCAO': 'com pendências nos controles técnicos'}.get(gate, gate)
    live_label = {'was_live': 'transmissão encerrada', 'not_live': 'vídeo publicado', 'is_live': 'ao vivo', 'is_upcoming': 'agendada', 'post_live': 'replay em processamento'}.get(status, status or NA)
    banner = f'Coleta {gate_label}. Estado do vídeo: {live_label}.'
    ws.merge_cells('A4:L5')
    put(ws, 4, 1, banner).alignment = Alignment(wrap_text=True, vertical='center')
    ws['A4'].fill = PatternFill('solid', fgColor=LIGHT if gate == 'APROVADO_PRODUCAO' else AMBER)
    # Seis cards compatíveis com a referência fornecida.
    for i, idx in enumerate([0, 3, 5, 6, 7, 8]):
        row, col = 7 + (i // 3) * 5, 1 + (i % 3) * 4
        end = col + 3
        label, value, _, fmt = a['indicators'][idx]
        ws.merge_cells(start_row=row, start_column=col, end_row=row+1, end_column=end)
        c = put(ws, row, col, label)
        c.font = Font(name='Arial', size=11, color='FFFFFF', bold=True)
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.merge_cells(start_row=row+2, start_column=col, end_row=row+3, end_column=end)
        c = put(ws, row+2, col, value)
        c.font = Font(name='Arial', size=24 if value is not None else 13, color=NAVY, bold=True)
        c.fill = PatternFill('solid', fgColor=LIGHT)
        c.number_format = fmt
        c.alignment = Alignment(vertical='center', horizontal='center')
    details = [
        f'Canal: {text(md.get("canal")) or NA}. Publicação: {display_date(md.get("data_upload"))}. Duração: {int(a["duration"] // 3600)}h {int(a["duration"] % 3600 // 60)}min {int(a["duration"] % 60)}s.' if a['duration'] else f'Canal: {text(md.get("canal")) or NA}. Duração indisponível.',
        f'Dados acumulados até: {display_date(md.get("data_hora_coleta"), True)}. Publicação não é necessariamente a data da transmissão.',
        'Público externo: autores que não pertencem ao canal oficial nem às contas institucionais/equipe identificadas. Participantes contam IDs, não espectadores.',
        'Perguntas, elogios e menções operacionais são sinais automáticos. Não representam validação humana, satisfação ou incidentes confirmados.',
    ]
    for row, label in zip([18, 20, 22, 24], details):
        ws.merge_cells(start_row=row, start_column=1, end_row=row+1, end_column=12)
        c = put(ws, row, 1, label)
        c.font = Font(name='Arial', size=11, color='243746')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    if a['chat_available']:
        chart(ws, generated['Analise_chat'], 'Natureza das mensagens (automática)', 2, 2, 4, 4+len(CATEGORIES), 1, 'A28')
        chart(ws, generated['Analise_chat'], 'Mensagens por período do vídeo', 2, 2, 18, 18+len(PERIODS), 1, 'G28')
    ws.merge_cells('A45:L46')
    put(ws, 45, 1, 'Fonte dos gráficos: Analise_chat / Mensagens_analise. Pré e pós-evento permanecem na análise. Distribuição do chat não mede retenção.').alignment = Alignment(wrap_text=True)
    observations = []
    if a['bins'] and any(v[0] for v in a['bins'].values()):
        peak = max(a['bins'], key=lambda b: a['bins'][b][0])
        observations.append(f'Maior volume do público externo: {a["bins"][peak][0]} mensagens na janela [{peak}, {peak+5}) min. Em empate, primeira janela.')
    if not a['chat_available']:
        observations.append('Chat indisponível: não interpretar ausência de registros como ausência de participação.')
    if not a['reliable']:
        observations.append('Tempo do chat não confirmado. Picos e quartis indisponíveis; revisar Qualidade_Coleta.')
    if len(public) < 30:
        observations.append(f'Base pequena ou ausente: {len(public)} mensagens do público externo; percentuais exigem cautela (corte descritivo: 30).')
    if not validation_ok:
        observations.append('Há divergência na conciliação. Consultar Validacao_relatorio antes de usar os totais.')
    if status in {'is_live', 'is_upcoming', 'post_live'}:
        observations.append('Transmissão em andamento/agendada ou em processamento: resultado parcial; repetir após a disponibilização do replay.')
    observations.append('Próxima ação: revisar as mensagens sinalizadas usando os filtros de Mensagens_analise e os links para a base original.')
    for i, label in enumerate(observations):
        row = 48 + i*2
        ws.merge_cells(start_row=row, start_column=1, end_row=row+1, end_column=12)
        c = put(ws, row, 1, label)
        c.font = Font(name='Arial', size=11, color='243746')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    nav_row = 50 + len(observations) * 2
    ws.merge_cells(f'A{nav_row}:F{nav_row+1}')
    put(ws, nav_row, 1, 'Consultar mensagens e classificações').hyperlink = "#'Mensagens_analise'!A1"
    ws.merge_cells(f'G{nav_row}:L{nav_row+1}')
    put(ws, nav_row, 7, 'Ler a transcrição').hyperlink = "#'Leitura_transcricao'!A1"
    ws.print_area = f'A1:L{nav_row+1}'
    ws.page_setup.fitToHeight = 1
    from relatorio_modelo import aplicar_modelo
    aplicar_modelo(wb, a)
    return a


def adicionar_relatorio(path, source_name='Coleta atual'):
    path = Path(path)
    wb = load_workbook(path)
    try:
        result = build_report(wb, source_name)
        wb.save(path)
        return result
    finally:
        wb.close()
