"""Auditoria offline da métrica canônica de participantes em um Excel gerado pelo projeto."""
from __future__ import annotations
import argparse
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from etl_core import canonical_author_type, participant_identity_metrics


def _records(ws):
    rows=list(ws.iter_rows(values_only=True))
    if not rows: return []
    hdr=list(rows[0])
    return [dict(zip(hdr,r)) for r in rows[1:] if any(v is not None for v in r)]


def _metric(ws, label, label_col=1, value_col=2):
    for row in ws.iter_rows(values_only=True):
        if str(row[label_col-1] or '').strip()==label:
            return row[value_col-1]
    return None


def auditar(path: str | Path) -> dict:
    path=Path(path)
    wb=load_workbook(path,data_only=True,read_only=True)
    try:
        md={r.get('campo'):r.get('valor') for r in _records(wb['Metadados'])}
        cid=str(md.get('canal_id') or '').strip()
        chat=pd.DataFrame(_records(wb['Chat_tratado']))
        if not chat.empty:
            chat['autor_tipo_canonico']=[canonical_author_type(a,aid,cid) for a,aid in zip(chat.get('autor',''),chat.get('canal_autor_id',''))]
        ident=participant_identity_metrics(chat,cid)
        expected=ident['participantes_publico_ids']
        checks={
            'Qualidade_Coleta': _metric(wb['Qualidade_Coleta'],'chat_participantes_publico_unicos') if 'Qualidade_Coleta' in wb.sheetnames else None,
            'Resumo': _metric(wb['Resumo'],'Participantes do público identificados') if 'Resumo' in wb.sheetnames else None,
            'Indicadores': _metric(wb['Indicadores'],'Participantes identificáveis') if 'Indicadores' in wb.sheetnames else None,
            'Participação': _metric(wb['Participação'],'Participantes únicos no vídeo') if 'Participação' in wb.sheetnames else None,
        }
        if 'Resumo da transmissão' in wb.sheetnames:
            ws=wb['Resumo da transmissão']
            headers=[c.value for c in ws[1]]
            for label in ('Contas participantes no chat','Participantes do Chat','Participantes identificáveis'):
                if label in headers:
                    checks['Resumo da transmissão']=ws.cell(2,headers.index(label)+1).value
                    break
        diverg={k:v for k,v in checks.items() if v is not None and int(v)!=int(expected)}
        return {'arquivo':str(path),'esperado_publico':expected,'identidade':ident,'abas':checks,'divergencias':diverg,'ok':not diverg}
    finally:
        wb.close()


def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('arquivo');args=ap.parse_args()
    r=auditar(args.arquivo)
    print('Participantes públicos:',r['esperado_publico'])
    print('Identidade:',r['identidade'])
    for k,v in r['abas'].items(): print(f'{k}: {v}')
    if not r['ok']:
        print('DIVERGÊNCIAS:',r['divergencias'])
        return 1
    print('AUDITORIA DE IDENTIDADE: OK')
    return 0

if __name__=='__main__': raise SystemExit(main())
