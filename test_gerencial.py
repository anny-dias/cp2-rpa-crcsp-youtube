"""Testes offline da entrada, método, preservação e integração com a V7."""
import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from entrada_youtube import normalizar_link, solicitar_link
from etl_core import RAW_CHAT_COLS, RAW_TRANS_COLS, participant_identity_metrics, canonical_author_type
from gerar_relatorio import gerar
from relatorio_gerencial import analyze, build_report, classify, period, SHEETS
import youtube_palestra_etl as etl


def fixture():
    wb = Workbook()
    del wb[wb.sheetnames[0]]
    data = {
        'Metadados': [['campo', 'valor'], ['video_id', 'qSwIRLA0C2o'], ['titulo', '=TEXTO DA FONTE'],
                      ['canal_id', 'official'], ['duracao_segundos', 1200], ['visualizacoes', 0], ['curtidas', 0]],
        'Qualidade_Coleta': [['indicador', 'valor'], ['chat_offset_confiavel', 'SIM'], ['status_geral', 'APROVADO_PRODUCAO']],
        'Status_Fontes': [['fonte', 'status'], ['Chat ao vivo', 'OK']],
        'Transcricao': [RAW_TRANS_COLS],
        'Transcricao_tratada': [['inicio', 'fim', 'texto', 'fonte', 'chave_bloco']],
        'Chat_ao_vivo': [RAW_CHAT_COLS],
        'Chat_tratado': [['chave_evento', 'offset_segundos', 'tempo_video', 'timestamp_usec', 'tipo',
                          'autor', 'autor_tipo', 'canal_autor_id', 'mensagem', 'id_mensagem']],
    }
    for name, rows in data.items():
        s = wb.create_sheet(name)
        for row in rows:
            s.append(row)
    # Mesmo nome, IDs diferentes; mesmo ID, nomes diferentes; sem ID; canal; sistema.
    messages = [
        ('m1', -1, 'A', 'id1', 'Oi'),
        ('m2', 0, 'A', 'id2', 'Como acessar o material?'),
        ('m3', 300, 'Novo nome', 'id1', 'Obrigado'),
        ('m4', 600, 'Sem ID', '', '=1+1'),
        ('m5', 1200, 'Canal', 'official', 'Bom dia'),
        ('m6', 1201, 'CRCSP falso', 'id3', 'sem som'),
    ]
    for mid, offset, name, aid, msg in messages:
        wb['Chat_ao_vivo'].append([offset, '', 100, 'mensagem', name, aid, msg, None, None, mid])
        wb['Chat_tratado'].append([mid, offset, '', 100, 'mensagem', name, 'canal_oficial', aid, msg, mid])
    # Duplicidade e sistema só no bruto.
    wb['Chat_ao_vivo'].append([0, '', 100, 'mensagem', 'A', 'id2', 'Como acessar o material?', None, None, 'm2'])
    wb['Chat_ao_vivo'].append([0, '', 101, 'sistema', '', '', 'Sistema', None, None, 's1'])
    return wb


class EntradaTest(unittest.TestCase):
    def test_formats(self):
        for url in ['https://www.youtube.com/watch?v=qSwIRLA0C2o&list=abc&t=50',
                    'youtu.be/qSwIRLA0C2o?si=abc', 'https://youtube.com/live/qSwIRLA0C2o',
                    'https://m.youtube.com/shorts/qSwIRLA0C2o', 'https://www.youtube-nocookie.com/embed/qSwIRLA0C2o']:
            self.assertEqual(normalizar_link(url), ('https://www.youtube.com/watch?v=qSwIRLA0C2o', 'qSwIRLA0C2o'))

    def test_invalid(self):
        for url in ['', 'http://evil.test/watch?v=qSwIRLA0C2o',
                    'https://youtube.com.evil.test/watch?v=qSwIRLA0C2o',
                    'https://youtube.com/@canal', 'https://youtube.com/playlist?list=abc',
                    'https://youtube.com/watch?v=bad', 'file:///youtube.com/watch?v=qSwIRLA0C2o',
                    'https://youtube.com/watch?v=qSwIRLA0C2o&v=abcdefghijk',
                    'https://user@youtube.com/watch?v=qSwIRLA0C2o']:
            with self.assertRaises(ValueError, msg=url):
                normalizar_link(url)

    def test_prompt_retries(self):
        with patch('builtins.input', side_effect=['bad', 'youtu.be/qSwIRLA0C2o']), contextlib.redirect_stdout(io.StringIO()):
            self.assertIn('qSwIRLA0C2o', solicitar_link())


class RelatorioTest(unittest.TestCase):
    def test_identity_population_and_totals(self):
        a = analyze(fixture())
        self.assertEqual(len(a['participants']), 2)
        self.assertEqual(len(a['public']), 4)
        self.assertEqual(len(a['official']), 1)
        self.assertEqual(len(a['institutional']), 1)
        self.assertEqual(len(a['participants']['id1']), 2)
        self.assertEqual(sum(a['counts'].values()), 4)
        self.assertEqual(sum(a['phases'].values()), 4)
        self.assertEqual(sum(v[0] for v in a['bins'].values()), 4)

    def test_canonical_identity_separates_channel_team_and_public(self):
        self.assertEqual(canonical_author_type('Canal', 'official', 'official'), 'canal_oficial')
        self.assertEqual(canonical_author_type('@Equipe_DES_CRCSP', 'team', 'official'), 'equipe_crcsp')
        self.assertEqual(canonical_author_type('Pessoa', 'p1', 'official'), 'publico_externo')
        df = pd.DataFrame([
            {'autor':'Canal','canal_autor_id':'official','autor_tipo':'canal_oficial'},
            {'autor':'@Equipe_DES_CRCSP','canal_autor_id':'team','autor_tipo':'equipe_crcsp'},
            {'autor':'Pessoa 1','canal_autor_id':'p1','autor_tipo':'publico_externo'},
            {'autor':'Pessoa 2','canal_autor_id':'p2','autor_tipo':'publico_externo'},
            {'autor':'Sem ID','canal_autor_id':'','autor_tipo':'sem_autor'},
        ])
        m = participant_identity_metrics(df, 'official')
        self.assertEqual(m['autores_identificados_total'], 4)
        self.assertEqual(m['institucionais_ids'], 2)
        self.assertEqual(m['participantes_publico_ids'], 2)
        self.assertEqual(m['autores_identificados_total'], m['institucionais_ids'] + m['participantes_publico_ids'])

    def test_time_boundaries(self):
        for t, expected in [(-1, 'Pré-evento'), (0, 'Q1'), (299.99, 'Q1'), (300, 'Q2'),
                            (600, 'Q3'), (900, 'Q4'), (1200, 'Q4'), (1201, 'Pós-evento'), (None, 'Indeterminado')]:
            self.assertEqual(period(t, 1200, True), expected)
        self.assertEqual(period(0, 1200, False), 'Indeterminado')
        self.assertEqual(period(0, None, True), 'Indeterminado')

    def test_automatic_not_fact(self):
        category, question, operation = classify('Onde encontro o certificado?')
        self.assertEqual(category, 'Certificado/EPC')
        self.assertTrue(question and operation)
        self.assertEqual(classify('O áudio está ótimo')[0], 'Áudio/conexão/imagem')
        self.assertTrue(classify('Gostaria de saber se posso usar esse crédito')[1])
        self.assertTrue(classify('É possível fazer dessa forma')[1])
        self.assertTrue(classify('Devo contabilizar assim')[1])
        self.assertFalse(classify('Excelente explicação sobre auditoria')[1])
        # Até elogio ao áudio é uma menção, não falha técnica confirmada.

    def test_unknown_chat_and_zero_views(self):
        wb = fixture()
        for name in ['Chat_tratado', 'Chat_ao_vivo']:
            wb[name].delete_rows(2, wb[name].max_row)
        wb['Status_Fontes']['B2'] = 'FALHOU'
        a = analyze(wb)
        metrics = {r[0]: r[1] for r in a['indicators']}
        self.assertEqual(metrics['Visualizações acumuladas'], 0)
        self.assertIsNone(metrics['Mensagens do público externo'])
        self.assertIsNone(metrics['Proporção de curtidas por visualização'])
        build_report(wb)

    def test_unreliable_time(self):
        wb = fixture()
        wb['Qualidade_Coleta']['B2'] = 'NAO'
        a = analyze(wb)
        self.assertFalse(a['bins'])
        self.assertEqual(a['phases']['Indeterminado'], 4)

    def test_negative_metadata(self):
        wb = fixture()
        wb['Metadados']['B6'] = -2
        metrics = {r[0]: r[1] for r in analyze(wb)['indicators']}
        self.assertIsNone(metrics['Visualizações acumuladas'])
        self.assertIsNone(metrics['Proporção de curtidas por visualização'])
        build_report(wb)
        rows = list(wb['Validacao_relatorio'].values)
        negative = next(r for r in rows if r[0] == 'Metadados numéricos negativos')
        self.assertEqual(negative[3], 'DIVERGÊNCIA')

    def test_preserve_sources_and_literal_strings(self):
        wb = fixture()
        before = {s.title: list(s.values) for s in wb}
        build_report(wb)
        for name, values in before.items():
            self.assertEqual(list(wb[name].values), values)
        self.assertEqual(wb['Mensagens_analise']['C8'].value, '=1+1')
        self.assertEqual(wb['Mensagens_analise']['C8'].data_type, 's')
        self.assertTrue(all(wb['Validacao_relatorio'].cell(r, 4).value == 'OK' for r in range(5, 13)))
        build_report(wb)  # Atualização idempotente sem duplicar abas.
        self.assertEqual(len([s for s in wb.sheetnames if s in SHEETS]), len(SHEETS))

    def test_conflict_and_same_path(self):
        wb = fixture()
        wb.create_sheet('Painel')
        with self.assertRaises(ValueError):
            build_report(wb)
        with tempfile.TemporaryDirectory() as temp:
            src = Path(temp) / 'src.xlsx'
            fixture().save(src)
            with self.assertRaises(ValueError):
                gerar(src, src)

    def test_offline_save_and_reopen(self):
        with tempfile.TemporaryDirectory() as temp:
            src = Path(temp) / 'src.xlsx'
            fixture().save(src)
            original = src.read_bytes()
            out = gerar(src)
            self.assertEqual(src.read_bytes(), original)
            wb = load_workbook(out)
            self.assertEqual(wb.sheetnames[0], 'Dashboard')
            self.assertEqual(wb.active.title, 'Dashboard')
            self.assertEqual(wb['Chat_ao_vivo'].sheet_state, 'visible')
            self.assertEqual(len(wb['Painel']._charts), 2)
            self.assertIn('Mensagens_analise', wb)
            wb.close()

    def test_pipeline_with_mocked_collectors(self):
        tr = pd.DataFrame([[0, 40, '00:00:00', '00:00:40', 'Fala', 'youtube']], columns=RAW_TRANS_COLS)
        chat = pd.DataFrame([[1, '00:00:01', 1000000, 'mensagem', 'A', 'id1', '=1+1', '', '', 'm1'],
                             [2, '00:00:02', 2000000, 'mensagem', 'B', 'id2', 'Olá', '', '', 'm2']], columns=RAW_CHAT_COLS)
        info = {'id': 'qSwIRLA0C2o', 'title': 'Teste', 'duration': 40, 'was_live': True, 'channel_id': 'official'}
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / 'report.xlsx'
            with patch.object(etl, '__file__', str(Path(temp) / 'youtube_palestra_etl.py')), \
                 patch.object(etl, 'metadata', return_value=info), \
                 patch.object(etl, 'collect_transcript', return_value=(tr, 'teste.json3', 'rede')), \
                 patch.object(etl, 'collect_chat', return_value=(chat, 'rede')), \
                 patch('sys.argv', ['etl', 'https://youtu.be/qSwIRLA0C2o', '--saida', str(out)]), \
                 contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(etl.main(), 0)
            wb = load_workbook(out)
            self.assertEqual(wb['Chat_ao_vivo']['G2'].data_type, 's')
            self.assertIn('Painel', wb)
            self.assertEqual(wb['Metadados']['B3'].value, 'https://www.youtube.com/watch?v=qSwIRLA0C2o')
            wb.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
